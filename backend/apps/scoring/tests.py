from unittest.mock import patch

import numpy as np
from django.test import TestCase

from apps.answers.models import Jawaban
from apps.exams.models import Ujian, UjianSoal
from apps.knowledge_base.models import KbChunk, KnowledgeBase, ProcessingStatus
from apps.master_data.models import MasterGuru, MasterJurusan, MasterKelas, MasterSiswa
from apps.questions.models import ConceptUnit, Soal, SoalKnowledgeBase
from apps.scoring.services.calibration import calibrate_semantic_score
from apps.scoring.services.pipeline import score_jawaban
from apps.scoring.services.text_quality import evaluate_text_quality
from apps.users.models import Role, User

TEXT_PARAMS = {
    "TEXT_MIN_WORDS": 5,
    "TEXT_MIN_ALPHA_RATIO": 0.75,
    "TEXT_MIN_VOWEL_RATIO": 0.50,
    "TEXT_MAX_UPPER_RATIO": 0.50,
}

SCORING_PARAMS = {
    "SEMANTIC_WEIGHT": 0.80,
    "CONCEPT_WEIGHT": 0.20,
    "TOP_K": 3,
    "CONCEPT_THRESHOLD": 0.60,
    "SEMANTIC_MIN_THRESHOLD": 0.25,
    "SEMANTIC_BASELINE": 0.40,
    "SEMANTIC_BELOW_BASELINE_MAX": 20,
    "NOISE_PENALTY": 0.10,
    **TEXT_PARAMS,
}


class CalibrationTests(TestCase):
    """Bagian 19: verifikasi rumus kalibrasi persis sesuai spesifikasi."""

    def test_di_bawah_threshold_nol(self):
        self.assertEqual(calibrate_semantic_score(0.10, t=0.25, b=0.40, m=20), 0.0)

    def test_tepat_di_threshold_nol(self):
        self.assertEqual(calibrate_semantic_score(0.25, t=0.25, b=0.40, m=20), 0.0)

    def test_antara_threshold_dan_baseline_linear(self):
        # s tepat di tengah t..b -> setengah dari m
        hasil = calibrate_semantic_score(0.325, t=0.25, b=0.40, m=20)
        self.assertAlmostEqual(hasil, 10.0, places=2)

    def test_tepat_di_baseline(self):
        self.assertAlmostEqual(calibrate_semantic_score(0.40, t=0.25, b=0.40, m=20), 20.0, places=2)

    def test_di_atas_baseline(self):
        # s=0.70 -> m + (100-m)*(0.70-0.40)/(1-0.40) = 20 + 80*0.5 = 60
        self.assertAlmostEqual(calibrate_semantic_score(0.70, t=0.25, b=0.40, m=20), 60.0, places=2)

    def test_maksimum_seratus(self):
        self.assertAlmostEqual(calibrate_semantic_score(1.0, t=0.25, b=0.40, m=20), 100.0, places=2)


class TextQualityTests(TestCase):
    """Bagian 18: singkatan seperti CPU/RAM/SQL tidak otomatis jadi noise."""

    def test_jawaban_normal_lolos(self):
        hasil = evaluate_text_quality(
            "Fotosintesis adalah proses pembuatan makanan oleh tumbuhan menggunakan cahaya matahari.",
            TEXT_PARAMS,
        )
        self.assertTrue(hasil.lolos)

    def test_singkatan_teknis_tidak_otomatis_gagal(self):
        hasil = evaluate_text_quality(
            "CPU dan RAM adalah komponen penting dalam sebuah komputer modern saat ini.",
            TEXT_PARAMS,
        )
        self.assertTrue(hasil.lolos)

    def test_jawaban_terlalu_pendek_gagal(self):
        hasil = evaluate_text_quality("ya benar", TEXT_PARAMS)
        self.assertFalse(hasil.lolos)

    def test_jawaban_full_kapital_gagal(self):
        hasil = evaluate_text_quality("INI JAWABAN SAYA YANG DITULIS SELURUHNYA DENGAN HURUF KAPITAL", TEXT_PARAMS)
        self.assertFalse(hasil.lolos)

    def test_jawaban_kosong_setelah_strip(self):
        hasil = evaluate_text_quality("   ", TEXT_PARAMS)
        self.assertEqual(hasil.jumlah_kata, 0)
        self.assertFalse(hasil.lolos)


def _fake_embed_texts(texts):
    """
    Stub deterministik pengganti SBERT (tidak butuh download model/internet).
    Vektor dibuat dari hash kata supaya teks yang identik/mirip punya
    similarity tinggi, dan teks acak punya similarity rendah -- cukup untuk
    memvalidasi ALUR pipeline (bukan kualitas semantik model asli).
    """
    import hashlib

    import numpy as np

    vectors = []
    for t in texts:
        words = t.lower().split()
        vec = np.zeros(64)
        for w in words:
            idx = int(hashlib.md5(w.encode()).hexdigest(), 16) % 64
            vec[idx] += 1
        norm = np.linalg.norm(vec)
        vectors.append((vec / norm).tolist() if norm else vec.tolist())
    return vectors


class ScoringPipelineTests(TestCase):
    """Bagian 34: uji end-to-end pipeline penilaian dengan embedder di-stub."""

    def setUp(self):
        kelas = MasterKelas.objects.create(nama_kelas="X TKJ 1", tingkat="X")
        jurusan = MasterJurusan.objects.create(kode_jurusan="TKJ", nama_jurusan="Teknik Komputer Jaringan")

        guru_user = User.objects.create_user(email="guru@sekolah.id", password="x", nama="Guru", role=Role.GURU)
        self.guru = MasterGuru.objects.create(user=guru_user)

        siswa_user = User.objects.create_user(email="siswa@sekolah.id", password=None, nama="Siswa", role=Role.SISWA)
        self.siswa = MasterSiswa.objects.create(user=siswa_user, nisn="001", kelas=kelas, jurusan=jurusan)

        self.kb = KnowledgeBase.objects.create(
            judul="Materi Fotosintesis", sumber="text",
            teks="fotosintesis adalah proses tumbuhan membuat makanan menggunakan cahaya matahari",
            guru=self.guru, processing_status=ProcessingStatus.DONE,
        )
        with patch("apps.knowledge_base.services.embedding.embed_texts", side_effect=_fake_embed_texts):
            from apps.knowledge_base.services.chunking import split_into_chunks
            from apps.knowledge_base.services.embedding import embed_texts

            chunk_texts = split_into_chunks(self.kb.teks)
            embeddings = embed_texts(chunk_texts)
            for i, (content, emb) in enumerate(zip(chunk_texts, embeddings)):
                KbChunk.objects.create(knowledge_base=self.kb, chunk_index=i, content=content, embedding=emb)

        self.soal = Soal.objects.create(guru=self.guru, pertanyaan="Jelaskan proses fotosintesis.")
        SoalKnowledgeBase.objects.create(soal=self.soal, knowledge_base=self.kb)
        self.cu = ConceptUnit.objects.create(soal=self.soal, konsep="cahaya matahari", urutan=1)

        self.ujian = Ujian.objects.create(
            nama_ujian="UTS", jenis_ujian="UTS", guru=self.guru, kelas=kelas, jurusan=jurusan, token="ABCDE"
        )
        UjianSoal.objects.create(ujian=self.ujian, soal=self.soal, urutan=1)

    @patch("apps.knowledge_base.services.embedding.embed_texts", side_effect=_fake_embed_texts)
    def test_jawaban_kosong_skor_nol(self, mock_embed):
        jawaban = Jawaban.objects.create(ujian=self.ujian, siswa=self.siswa, soal=self.soal, jawaban_teks="")
        penilaian = score_jawaban(jawaban)
        self.assertEqual(penilaian.final_score, 0.0)
        self.assertEqual(penilaian.relevance_status, "gagal")
        self.assertEqual(penilaian.processing_status, "success")

    @patch("apps.knowledge_base.services.embedding.embed_texts", side_effect=_fake_embed_texts)
    def test_jawaban_relevan_dapat_skor_di_atas_nol(self, mock_embed):
        jawaban = Jawaban.objects.create(
            ujian=self.ujian, siswa=self.siswa, soal=self.soal,
            jawaban_teks="fotosintesis adalah proses tumbuhan membuat makanan menggunakan cahaya matahari",
        )
        penilaian = score_jawaban(jawaban)
        self.assertIsNotNone(penilaian.final_score)
        self.assertGreater(penilaian.final_score, 0.0)
        self.assertEqual(penilaian.relevance_status, "lolos")

    @patch("apps.knowledge_base.services.embedding.embed_texts", side_effect=_fake_embed_texts)
    def test_jawaban_tidak_relevan_skor_nol(self, mock_embed):
        jawaban = Jawaban.objects.create(
            ujian=self.ujian, siswa=self.siswa, soal=self.soal,
            jawaban_teks="kucing saya suka bermain bola di halaman setiap sore hari",
        )
        penilaian = score_jawaban(jawaban)
        self.assertEqual(penilaian.final_score, 0.0)
        self.assertEqual(penilaian.relevance_status, "gagal")

    def test_jawaban_ganda_ditolak_database(self):
        from django.db import IntegrityError

        Jawaban.objects.create(ujian=self.ujian, siswa=self.siswa, soal=self.soal, jawaban_teks="jawaban pertama")
        with self.assertRaises(IntegrityError):
            Jawaban.objects.create(ujian=self.ujian, siswa=self.siswa, soal=self.soal, jawaban_teks="jawaban kedua")


class ConceptSplittingRegressionTests(TestCase):
    """
    Regresi kasus nyata yang Anda laporkan: jawaban esai yang mendefinisikan
    beberapa Concept Unit dalam SATU kalimat majemuk (dipisah koma +
    "sedangkan") harus tetap bisa terdeteksi per-konsep, bukan cuma 2 dari 5.
    """

    def setUp(self):
        self.cu_tesis = type("CU", (), {"id": 3, "konsep": "Tesis - bagian yang berisi pengenalan topik atau pendapat awal penulis"})()
        self.cu_argumentasi = type("CU", (), {"id": 4, "konsep": "Argumentasi - bagian yang berisi alasan, penjelasan, fakta, atau data yang mendukung tesis"})()
        self.cu_penegasan = type("CU", (), {"id": 5, "konsep": "Penegasan ulang - bagian yang menegaskan kembali pendapat atau gagasan utama yang telah disampaikan"})()

        self.jawaban = (
            "Struktur teks eksposisi terdiri atas tesis, argumentasi, dan penegasan ulang. "
            "Tesis berisi pengenalan topik atau pendapat awal penulis, argumentasi berisi alasan, "
            "penjelasan, fakta, atau data yang mendukung tesis, sedangkan penegasan ulang berisi "
            "penegasan kembali pendapat atau gagasan utama yang telah disampaikan."
        )

    @patch("apps.knowledge_base.services.embedding.embed_texts", side_effect=_fake_embed_texts)
    def test_klausa_tesis_kini_jadi_kandidat_tersendiri(self, mock_embed):
        from apps.scoring.services.concept_scoring import _split_kalimat

        kandidat = _split_kalimat(self.jawaban)
        # Klausa bersih "Tesis berisi ..." harus muncul sebagai kandidat
        # tersendiri (sebelum perbaikan, ini hanya ada sebagai bagian dari
        # satu kalimat majemuk yang mengandung 3 definisi konsep sekaligus).
        self.assertTrue(
            any(k.startswith("Tesis berisi pengenalan topik") for k in kandidat),
            f"Klausa Tesis tidak ditemukan sebagai kandidat terpisah: {kandidat}",
        )
        self.assertTrue(
            any(k.startswith("sedangkan penegasan ulang berisi") for k in kandidat),
            f"Klausa penegasan ulang tidak ditemukan sebagai kandidat terpisah: {kandidat}",
        )

    @patch("apps.knowledge_base.services.embedding.embed_texts", side_effect=_fake_embed_texts)
    def test_similarity_tesis_naik_dibanding_hanya_kalimat_utuh(self, mock_embed):
        """
        Bandingkan similarity CU Tesis: pakai klausa vs pakai kalimat utuh
        yang tercampur 3 konsep. Similarity dgn klausa harus >= similarity
        dgn kalimat penuh -- inti dari perbaikan ini (tidak pernah menurun).
        """
        from apps.scoring.services.concept_scoring import _cosine_sim

        klausa_bersih = "Tesis berisi pengenalan topik atau pendapat awal penulis"
        kalimat_campur = (
            "Tesis berisi pengenalan topik atau pendapat awal penulis, argumentasi berisi alasan, "
            "penjelasan, fakta, atau data yang mendukung tesis, sedangkan penegasan ulang berisi "
            "penegasan kembali pendapat atau gagasan utama yang telah disampaikan."
        )
        cu_vec = np.array(_fake_embed_texts([self.cu_tesis.konsep])[0])
        sim_klausa = _cosine_sim(cu_vec, np.array(_fake_embed_texts([klausa_bersih])[0]))
        sim_campur = _cosine_sim(cu_vec, np.array(_fake_embed_texts([kalimat_campur])[0]))

        self.assertGreaterEqual(sim_klausa, sim_campur)


class TopKChunksPersistenceTests(TestCase):
    """Chunk relevan (top-K) harus tersimpan di Penilaian, atas permintaan Anda, untuk ditampilkan di detail hasil."""

    def setUp(self):
        kelas = MasterKelas.objects.create(nama_kelas="X TKJ 3", tingkat="X")
        jurusan = MasterJurusan.objects.create(kode_jurusan="TKJ3", nama_jurusan="TKJ 3")
        guru_user = User.objects.create_user(email="topk-guru@sekolah.id", password="x", nama="Guru", role=Role.GURU)
        guru = MasterGuru.objects.create(user=guru_user)
        siswa_user = User.objects.create_user(email="topk-siswa@sekolah.id", password=None, nama="Siswa", role=Role.SISWA)
        siswa = MasterSiswa.objects.create(user=siswa_user, nisn="0011", kelas=kelas, jurusan=jurusan)

        kb = KnowledgeBase.objects.create(
            judul="Materi Top-K", sumber="text", teks="materi relevan",
            guru=guru, processing_status=ProcessingStatus.DONE,
        )
        with patch("apps.knowledge_base.services.embedding.embed_texts", side_effect=_fake_embed_texts):
            from apps.knowledge_base.services.chunking import split_into_chunks
            from apps.knowledge_base.services.embedding import embed_texts
            chunk_texts = split_into_chunks(kb.teks)
            embeddings = embed_texts(chunk_texts)
            for i, (content, emb) in enumerate(zip(chunk_texts, embeddings)):
                KbChunk.objects.create(knowledge_base=kb, chunk_index=i, content=content, embedding=emb)

        soal = Soal.objects.create(guru=guru, pertanyaan="Jelaskan materi relevan.")
        SoalKnowledgeBase.objects.create(soal=soal, knowledge_base=kb)
        ConceptUnit.objects.create(soal=soal, konsep="materi relevan", urutan=1)

        ujian = Ujian.objects.create(nama_ujian="UH", jenis_ujian="Ulangan Harian", guru=guru, kelas=kelas, jurusan=jurusan, token="TOPKX")
        UjianSoal.objects.create(ujian=ujian, soal=soal, urutan=1)
        self.jawaban = Jawaban.objects.create(ujian=ujian, siswa=siswa, soal=soal, jawaban_teks="materi relevan sekali")

    @patch("apps.knowledge_base.services.embedding.embed_texts", side_effect=_fake_embed_texts)
    def test_top_k_chunks_tersimpan_dengan_isi_dan_similarity(self, mock_embed):
        penilaian = score_jawaban(self.jawaban)
        self.assertTrue(len(penilaian.top_k_chunks) > 0)
        chunk_info = penilaian.top_k_chunks[0]
        self.assertIn("chunk_id", chunk_info)
        self.assertIn("similarity", chunk_info)
        self.assertIn("content", chunk_info)
        self.assertIn("knowledge_base_judul", chunk_info)
        self.assertEqual(chunk_info["knowledge_base_judul"], "Materi Top-K")

    @patch("apps.knowledge_base.services.embedding.embed_texts", side_effect=_fake_embed_texts)
    def test_top_k_chunks_kosong_untuk_jawaban_kosong(self, mock_embed):
        self.jawaban.jawaban_teks = ""
        self.jawaban.save()
        penilaian = score_jawaban(self.jawaban)
        self.assertEqual(penilaian.top_k_chunks, [])


class ExportLaporanTests(TestCase):
    """Ekspor = laporan profesional: info ujian dari DB + rekap nilai akhir per siswa (bagian 21 & 29)."""

    def setUp(self):
        from rest_framework.test import APIClient

        from apps.scoring.models import Penilaian

        kelas = MasterKelas.objects.create(nama_kelas="XI RPL 2", tingkat="XI")
        jurusan = MasterJurusan.objects.create(kode_jurusan="RPL", nama_jurusan="Rekayasa Perangkat Lunak")
        guru_user = User.objects.create_user(email="lap-guru@sekolah.id", password="Rahasia1", nama="Ibu Sari Wulandari", role=Role.GURU)
        guru = MasterGuru.objects.create(user=guru_user)

        soal1 = Soal.objects.create(guru=guru, pertanyaan="Jelaskan teks eksposisi.")
        soal2 = Soal.objects.create(guru=guru, pertanyaan="Jelaskan teks narasi.")
        ujian = Ujian.objects.create(
            nama_ujian="UTS Bahasa Indonesia", jenis_ujian="UTS", guru=guru,
            kelas=kelas, jurusan=jurusan, token="LAPOR",
        )
        UjianSoal.objects.create(ujian=ujian, soal=soal1, urutan=1)
        UjianSoal.objects.create(ujian=ujian, soal=soal2, urutan=2)

        def buat_siswa(nama, nisn):
            u = User.objects.create_user(email=f"{nisn}@sekolah.id", password=None, nama=nama, role=Role.SISWA)
            return MasterSiswa.objects.create(user=u, nisn=nisn, kelas=kelas, jurusan=jurusan)

        def nilai(siswa, soal, skor, status="success"):
            j = Jawaban.objects.create(ujian=ujian, siswa=siswa, soal=soal, jawaban_teks="jawaban")
            Penilaian.objects.create(jawaban=j, final_score=skor, processing_status=status)

        andi = buat_siswa("Andi Pratama", "9001")   # 80 & 60 -> 70.00
        budi = buat_siswa("Budi Hartono", "9002")   # hanya 1 soal (90) -> 90/2 = 45.00
        nilai(andi, soal1, 80.0)
        nilai(andi, soal2, 60.0)
        nilai(budi, soal1, 90.0)

        self.client = APIClient()
        token = self.client.post("/api/auth/login/", {"email": "lap-guru@sekolah.id", "credential": "Rahasia1"}, format="json").json()["access"]
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {token}")

    def _sel_xlsx(self):
        from io import BytesIO

        from openpyxl import load_workbook

        res = self.client.get("/api/penilaian/export/xlsx/")
        self.assertEqual(res.status_code, 200)
        wb = load_workbook(BytesIO(res.content))
        ws = wb["Rekap Nilai"]
        return wb, [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]

    def test_excel_memuat_blok_info_dari_database(self):
        _, sel = self._sel_xlsx()
        teks = " | ".join(sel)
        self.assertIn("UTS Bahasa Indonesia", teks)
        self.assertIn("XI RPL 2", teks)
        self.assertIn("Rekayasa Perangkat Lunak", teks)
        self.assertIn("Ibu Sari Wulandari", teks)
        self.assertIn("Tanggal Pelaksanaan", teks)

    def test_excel_nilai_akhir_rata_rata_seluruh_soal_belum_dijawab_nol(self):
        wb, _ = self._sel_xlsx()
        ws = wb["Rekap Nilai"]
        nilai_per_nama = {}
        for row in ws.iter_rows(values_only=True):
            if row and row[1] in ("Andi Pratama", "Budi Hartono"):
                nilai_per_nama[row[1]] = row[6]
        self.assertEqual(nilai_per_nama["Andi Pratama"], 70.0)
        self.assertEqual(nilai_per_nama["Budi Hartono"], 45.0)

    def test_excel_keterangan_soal_belum_dijawab(self):
        _, sel = self._sel_xlsx()
        self.assertIn("1 soal belum dijawab", " | ".join(sel))

    def test_excel_punya_sheet_rincian_per_soal(self):
        wb, _ = self._sel_xlsx()
        self.assertIn("Rincian Per Soal", wb.sheetnames)
        ws = wb["Rincian Per Soal"]
        semua = " | ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        self.assertIn("Jelaskan teks eksposisi.", semua)

    def test_soal_gagal_diproses_tidak_dipaksa_jadi_nol(self):
        from apps.scoring.models import Penilaian

        p = Penilaian.objects.filter(jawaban__siswa__nisn="9001").first()
        p.final_score = None
        p.processing_status = "failed"
        p.save()
        wb, sel = self._sel_xlsx()
        ws = wb["Rekap Nilai"]
        andi = [row for row in ws.iter_rows(values_only=True) if row and row[1] == "Andi Pratama"][0]
        self.assertEqual(andi[6], "-")
        self.assertIn("Ada soal gagal diproses", " | ".join(sel))

    def test_pdf_valid_dan_memuat_info_ujian(self):
        from io import BytesIO

        from PyPDF2 import PdfReader

        res = self.client.get("/api/penilaian/export/pdf/")
        self.assertEqual(res.status_code, 200)
        self.assertTrue(res.content.startswith(b"%PDF"))
        teks = " ".join(page.extract_text() or "" for page in PdfReader(BytesIO(res.content)).pages)
        self.assertIn("UTS Bahasa Indonesia", teks)
        self.assertIn("Ibu Sari Wulandari", teks)
        self.assertIn("Andi Pratama", teks)
        self.assertIn("70.00", teks)

    def test_nama_dengan_karakter_khusus_tidak_merusak_pdf(self):
        u = User.objects.create_user(email="khusus@sekolah.id", password=None, nama="Siti <b>& Co", role=Role.SISWA)
        MasterSiswa.objects.create(
            user=u, nisn="9003",
            kelas=MasterKelas.objects.get(nama_kelas="XI RPL 2"),
            jurusan=MasterJurusan.objects.get(kode_jurusan="RPL"),
        )
        from apps.scoring.models import Penilaian

        ujian = Ujian.objects.get(token="LAPOR")
        j = Jawaban.objects.create(ujian=ujian, siswa=MasterSiswa.objects.get(nisn="9003"), soal=Soal.objects.first(), jawaban_teks="x")
        Penilaian.objects.create(jawaban=j, final_score=50.0, processing_status="success")
        res = self.client.get("/api/penilaian/export/pdf/")
        self.assertEqual(res.status_code, 200)

    def test_export_kosong_tidak_error(self):
        res_x = self.client.get("/api/penilaian/export/xlsx/?search=tidakadanamaini")
        res_p = self.client.get("/api/penilaian/export/pdf/?search=tidakadanamaini")
        self.assertEqual(res_x.status_code, 200)
        self.assertEqual(res_p.status_code, 200)