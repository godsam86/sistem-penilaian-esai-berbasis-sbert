import io

from django.db import transaction
from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rest_framework import viewsets
from rest_framework.decorators import action
from rest_framework.filters import SearchFilter
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from apps.activity_logs.utils import log_activity
from apps.master_data.models import MasterAdmin, MasterGuru, MasterJurusan, MasterKelas, MasterSiswa
from apps.master_data.serializers import (
    AdminAccountSerializer,
    GuruAccountSerializer,
    MasterJurusanSerializer,
    MasterKelasSerializer,
    SiswaAccountSerializer,
    SiswaImportSerializer,
)
from apps.users.models import Role, Status, User
from apps.users.permissions import IsAdmin


def _nilai_sel(v):
    """Ubah isi sel Excel jadi teks bersih. Angka bulat (mis. 12345.0) tidak boleh jadi '12345.0'."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


class ReadableByAnyAuthenticatedMixin:
    """
    Guru butuh baca daftar kelas/jurusan untuk membuat ujian (bagian 8), tapi
    tidak boleh mengubahnya -- itu tetap wewenang admin (bagian 25).
    """

    def get_permissions(self):
        if self.action in ("list", "retrieve"):
            return [IsAuthenticated()]
        return [IsAdmin()]



class SoftDeactivateMixin:
    """
    Bagian 25: admin tidak menghapus, hanya menonaktifkan (status=0).
    Riwayat transaksi terkait tetap aman (bagian 28). Bisa diaktifkan
    kembali lewat action `aktifkan` (atas permintaan Anda).
    """

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        target = getattr(instance, "user", instance)  # akun -> nonaktifkan users.status
        target.status = 0
        target.save(update_fields=["status"])
        log_activity(
            request.user,
            action="nonaktifkan",
            module=self.activity_module,
            description=f"id={instance.pk}",
            request=request,
        )
        return Response(status=204)

    @action(detail=True, methods=["post"])
    def aktifkan(self, request, pk=None):
        instance = self.get_object()
        target = getattr(instance, "user", instance)
        target.status = 1
        target.save(update_fields=["status"])
        log_activity(
            request.user,
            action="aktifkan",
            module=self.activity_module,
            description=f"id={instance.pk}",
            request=request,
        )
        serializer = self.get_serializer(instance)
        return Response(serializer.data)


class MasterKelasViewSet(ReadableByAnyAuthenticatedMixin, viewsets.ModelViewSet):
    queryset = MasterKelas.objects.all().order_by("nama_kelas")
    serializer_class = MasterKelasSerializer
    permission_classes = [IsAdmin]
    filter_backends = [SearchFilter]
    search_fields = ["nama_kelas", "tingkat"]
    activity_module = "master_kelas"
    # Data referensi untuk dropdown (bagian 5) -- selalu lengkap, tidak
    # dipaginasi seperti daftar akun yang bisa sangat banyak.
    pagination_class = None

    def perform_create(self, serializer):
        instance = serializer.save(created_by=self.request.user.id)
        log_activity(self.request.user, "buat", self.activity_module, str(instance.pk), self.request)

    def perform_update(self, serializer):
        instance = serializer.save(updated_by=self.request.user.id)
        log_activity(self.request.user, "ubah", self.activity_module, str(instance.pk), self.request)


class MasterJurusanViewSet(ReadableByAnyAuthenticatedMixin, viewsets.ModelViewSet):
    queryset = MasterJurusan.objects.all().order_by("nama_jurusan")
    serializer_class = MasterJurusanSerializer
    permission_classes = [IsAdmin]
    filter_backends = [SearchFilter]
    search_fields = ["nama_jurusan", "kode_jurusan"]
    activity_module = "master_jurusan"
    pagination_class = None

    def perform_create(self, serializer):
        instance = serializer.save(created_by=self.request.user.id)
        log_activity(self.request.user, "buat", self.activity_module, str(instance.pk), self.request)

    def perform_update(self, serializer):
        instance = serializer.save(updated_by=self.request.user.id)
        log_activity(self.request.user, "ubah", self.activity_module, str(instance.pk), self.request)


class GuruAccountViewSet(SoftDeactivateMixin, viewsets.ModelViewSet):
    queryset = MasterGuru.objects.select_related("user").order_by("user__nama")
    serializer_class = GuruAccountSerializer
    permission_classes = [IsAdmin]
    filter_backends = [SearchFilter]
    search_fields = ["user__nama", "user__email", "nip"]
    activity_module = "master_guru"

    def perform_create(self, serializer):
        instance = serializer.save(created_by=self.request.user.id)
        log_activity(self.request.user, "buat_akun_guru", self.activity_module, str(instance.pk), self.request)

    def perform_update(self, serializer):
        instance = serializer.save(updated_by=self.request.user.id)
        log_activity(self.request.user, "ubah_akun_guru", self.activity_module, str(instance.pk), self.request)


class AdminAccountViewSet(SoftDeactivateMixin, viewsets.ModelViewSet):
    """
    Bagian 25 (perluasan atas permintaan Anda): admin mengelola akun admin
    lain lewat master_admin, pola sama seperti master_guru/master_siswa.
    """

    queryset = MasterAdmin.objects.select_related("user").order_by("user__nama")
    serializer_class = AdminAccountSerializer
    permission_classes = [IsAdmin]
    filter_backends = [SearchFilter]
    search_fields = ["user__nama", "user__email", "jabatan"]
    activity_module = "master_admin"

    def perform_create(self, serializer):
        instance = serializer.save(created_by=self.request.user.id)
        log_activity(self.request.user, "buat_akun_admin", self.activity_module, str(instance.pk), self.request)

    def perform_update(self, serializer):
        instance = serializer.save(updated_by=self.request.user.id)
        log_activity(self.request.user, "ubah_akun_admin", self.activity_module, str(instance.pk), self.request)

    def destroy(self, request, *args, **kwargs):
        # Cegah admin menonaktifkan akunnya sendiri -- supaya tidak ada yang
        # tidak sengaja mengunci diri sendiri keluar dari sistem.
        instance = self.get_object()
        if instance.user_id == request.user.id:
            return Response(
                {"detail": "Anda tidak dapat menonaktifkan akun Anda sendiri."}, status=400
            )
        return super().destroy(request, *args, **kwargs)


class SiswaAccountViewSet(SoftDeactivateMixin, viewsets.ModelViewSet):
    queryset = MasterSiswa.objects.select_related("user", "kelas", "jurusan").order_by("user__nama")
    serializer_class = SiswaAccountSerializer
    permission_classes = [IsAdmin]
    filter_backends = [SearchFilter]
    search_fields = ["user__nama", "nisn"]
    activity_module = "master_siswa"

    def get_queryset(self):
        qs = super().get_queryset()
        kelas_id = self.request.query_params.get("kelas_id")
        jurusan_id = self.request.query_params.get("jurusan_id")
        if kelas_id:
            qs = qs.filter(kelas_id=kelas_id)
        if jurusan_id:
            qs = qs.filter(jurusan_id=jurusan_id)
        return qs

    def perform_create(self, serializer):
        instance = serializer.save(created_by=self.request.user.id)
        log_activity(self.request.user, "buat_akun_siswa", self.activity_module, str(instance.pk), self.request)

    def perform_update(self, serializer):
        instance = serializer.save(updated_by=self.request.user.id)
        log_activity(self.request.user, "ubah_akun_siswa", self.activity_module, str(instance.pk), self.request)

    @action(detail=False, methods=["get"], url_path="template")
    def template(self, request):
        """Unduh template Excel kosong (kolom: Nama, Email, NISN) untuk impor massal."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Data Siswa"

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="1D4ED8")
        for idx, (judul, lebar) in enumerate([("Nama", 34), ("Email", 34), ("NISN", 18)], start=1):
            cell = ws.cell(row=1, column=idx, value=judul)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(idx)].width = lebar
        # Kolom NISN berformat TEKS supaya angka 0 di depan tidak hilang saat diketik.
        ws.column_dimensions["C"].number_format = "@"

        petunjuk = wb.create_sheet("Petunjuk")
        petunjuk.column_dimensions["A"].width = 90
        for baris in [
            "PETUNJUK PENGISIAN",
            "1. Isi data siswa di sheet 'Data Siswa', mulai dari baris ke-2 (baris 1 adalah judul kolom, jangan diubah).",
            "2. Kolom: Nama, Email, NISN. Email dan NISN tidak boleh sama dengan siswa lain yang sudah terdaftar.",
            "3. Semua siswa dalam satu file akan dimasukkan ke SATU kelas & jurusan yang dipilih saat impor.",
            "4. Contoh isi:  Budi Santoso | budi@sekolah.id | 0012345678",
            "5. Siswa login memakai Email + NISN.",
        ]:
            petunjuk.append([baris])
        petunjuk["A1"].font = Font(bold=True)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        response = HttpResponse(
            buffer.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = "attachment; filename=template-impor-siswa.xlsx"
        return response

    @action(detail=False, methods=["post"], url_path="impor", parser_classes=[MultiPartParser, FormParser])
    def impor(self, request):
        """
        Impor massal siswa dari Excel (kolom: Nama, Email, NISN) untuk SATU
        kelas & jurusan tujuan yang dipilih di form -- atas permintaan Anda.
        Baris yang gagal (email/NISN duplikat, kolom kosong, dst) dilaporkan
        satu per satu, baris lain yang valid tetap diproses (bukan all-or-nothing).
        """
        serializer = SiswaImportSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        file = serializer.validated_data["file"]
        kelas = serializer.validated_data["kelas"]
        jurusan = serializer.validated_data["jurusan"]

        try:
            wb = load_workbook(file, data_only=True)
            ws = wb.active
        except Exception:
            return Response({"detail": "File Excel tidak bisa dibaca. Pastikan formatnya .xlsx."}, status=400)

        berhasil = []
        gagal = []

        for baris_ke, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            nama, email, nisn = [_nilai_sel(v) for v in (list(row) + [None, None, None])[:3]]
            if not nama and not email and not nisn:
                continue  # baris kosong (mis. sisa format di ujung file) -- lewati tanpa dianggap error
            if not nama or not email or not nisn:
                gagal.append({"baris": baris_ke, "alasan": "Nama, Email, atau NISN kosong."})
                continue

            if User.objects.filter(email__iexact=email).exists():
                gagal.append({"baris": baris_ke, "alasan": f"Email '{email}' sudah dipakai akun lain."})
                continue
            if MasterSiswa.objects.filter(nisn=nisn).exists():
                gagal.append({"baris": baris_ke, "alasan": f"NISN '{nisn}' sudah dipakai siswa lain."})
                continue

            try:
                with transaction.atomic():
                    user = User.objects.create_user(
                        email=email, password=None, nama=nama, role=Role.SISWA, status=Status.AKTIF,
                    )
                    MasterSiswa.objects.create(
                        user=user, nisn=nisn, kelas=kelas, jurusan=jurusan, created_by=request.user.id,
                    )
                berhasil.append({"baris": baris_ke, "nama": nama, "nisn": nisn})
            except Exception as exc:
                gagal.append({"baris": baris_ke, "alasan": f"Gagal disimpan: {exc}"})

        log_activity(
            request.user, "impor_siswa", self.activity_module,
            f"berhasil={len(berhasil)} gagal={len(gagal)} kelas={kelas.id} jurusan={jurusan.id}", request,
        )
        return Response({"berhasil": berhasil, "gagal": gagal})